"""Core integration tests. Run with: python -m unittest discover -s tests -v"""
from __future__ import annotations

import os
import shutil
import tempfile
import time
import unittest
import zipfile
from pathlib import Path

TEST_DATA = Path(tempfile.mkdtemp(prefix="qiming_tests_"))
os.environ["QIMING_DATA_DIR"] = str(TEST_DATA)
os.environ["QIMING_MULTI"] = "1"
os.environ["QIMING_EMAIL_CAPTURE_DIR"] = str(TEST_DATA / "mail")

from backend import create_app  # noqa: E402
from backend.ai import local_outline  # noqa: E402
from backend.themes import get_theme  # noqa: E402


class QimingApiTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = create_app(TESTING=True)
        cls.client = cls.app.test_client()

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEST_DATA, ignore_errors=True)

    def test_01_invite_registration(self):
        bad = self.client.post(
            "/api/auth/register",
            json={"username": "tester", "password": "secret123", "invite_code": "wrong"},
        )
        self.assertEqual(bad.status_code, 400)
        good = self.client.post(
            "/api/auth/register",
            json={
                "username": "tester",
                "display_name": "测试用户",
                "email": "tester@example.com",
                "password": "secret123",
                "invite_code": "songyonghao2016",
            },
        )
        self.assertEqual(good.status_code, 200, good.get_json())
        self.assertTrue(good.get_json()["ok"])
        me = self.client.get("/api/auth/me").get_json()
        self.assertEqual(me["user"]["username"], "tester")

    def test_02_duplicate_and_invite_case(self):
        duplicate = self.client.post(
            "/api/auth/register",
            json={"username": "tester", "password": "secret123", "invite_code": "songyonghao2016"},
        )
        self.assertEqual(duplicate.status_code, 409)
        upper = self.client.post(
            "/api/auth/register",
            json={"username": "tester2", "password": "secret123", "invite_code": "SONGYONGHAO2016"},
        )
        self.assertEqual(upper.status_code, 400)

    def test_03_outline_project_generate(self):
        outline_response = self.client.post(
            "/api/ai/outline",
            json={"prompt": "人工智能如何改变教育行业", "slides": 7, "provider": "local"},
        )
        self.assertEqual(outline_response.status_code, 200)
        outline = outline_response.get_json()["outline"]
        self.assertEqual(len(outline["slides"]), 7)
        self.assertEqual(outline["slides"][0]["type"], "cover")

        create = self.client.post(
            "/api/projects",
            json={
                "title": outline["title"],
                "prompt": "人工智能如何改变教育行业",
                "outline": outline,
                "template_id": "ink-paper",
                "options": {"animations": True, "audio": False, "video": False, "auto_advance": False},
            },
        )
        self.assertEqual(create.status_code, 200, create.get_json())
        project = create.get_json()["project"]
        project_id = project["id"]

        start = self.client.post(f"/api/projects/{project_id}/generate")
        self.assertEqual(start.status_code, 200)
        state = {}
        for _ in range(40):
            time.sleep(0.6)
            state = self.client.get(f"/api/projects/{project_id}/status").get_json()
            if state["status"] in {"ready", "failed"}:
                break
        self.assertEqual(state.get("status"), "ready", state)

        download = self.client.get(f"/api/projects/{project_id}/download?format=pptx")
        self.assertEqual(download.status_code, 200)
        self.assertGreater(len(download.data), 100_000)
        self.assertTrue(download.data[:2] == b"PK")
        # Keep a copy on disk for structural assertions.
        target = TEST_DATA / "generated.pptx"
        target.write_bytes(download.data)
        download.close()
        with zipfile.ZipFile(target) as archive:
            names = archive.namelist()
            slides = [name for name in names if name.startswith("ppt/slides/slide") and name.endswith(".xml")]
            self.assertEqual(len(slides), 7)
            first = archive.read("ppt/slides/slide1.xml")
            self.assertIn(b"p:timing", first)
            self.assertIn(b"p:transition", first)
            self.assertTrue(any(name.endswith(".wav") for name in names) is False)

    def test_04_no_animation_deck(self):
        outline = local_outline("如何做一次有效复盘", target_slides=6)
        output = TEST_DATA / "noanim"
        output.mkdir(exist_ok=True)
        from backend.pptx_builder import build_presentation

        path, meta = build_presentation(
            project_dir=output,
            title=outline["title"],
            outline=outline,
            theme=get_theme("ink-paper"),
            options={"animations": False, "audio": False, "video": False, "auto_advance": False},
        )
        self.assertEqual(meta["validation"]["slides"], 6)
        self.assertEqual(meta["validation"]["animations"], 0)
        self.assertEqual(meta["validation"]["transitions"], 0)

    def test_05_admin_permissions_and_stats(self):
        user_client = self.app.test_client()
        login = user_client.post("/api/auth/login", json={"username": "tester", "password": "secret123"})
        self.assertEqual(login.status_code, 200)
        denied = user_client.get("/api/admin/overview")
        self.assertEqual(denied.status_code, 403)

        admin_client = self.app.test_client()
        admin_login = admin_client.post("/api/auth/login", json={"username": "admin", "password": "Admin@2016"})
        self.assertEqual(admin_login.status_code, 200)
        overview = admin_client.get("/api/admin/overview")
        self.assertEqual(overview.status_code, 200)
        payload = overview.get_json()
        self.assertGreaterEqual(payload["stats"]["total_users"], 1)
        self.assertEqual(payload["invite"]["code"], "songyonghao2016")
        self.assertGreaterEqual(payload["invite"]["uses"], 1)
        users = admin_client.get("/api/admin/users").get_json()
        self.assertTrue(any(item["username"] == "tester" for item in users["items"]))
        settings = admin_client.get("/api/admin/settings")
        self.assertEqual(settings.status_code, 200)

    def test_06_profile_and_password(self):
        client = self.app.test_client()
        client.post("/api/auth/login", json={"username": "tester", "password": "secret123"})
        update = client.patch(
            "/api/auth/profile",
            json={"display_name": "新昵称", "email": "new@example.com"},
        )
        self.assertEqual(update.status_code, 200)
        self.assertEqual(update.get_json()["user"]["display_name"], "新昵称")
        change = client.post(
            "/api/auth/password",
            json={"current_password": "secret123", "new_password": "newsecret123"},
        )
        self.assertEqual(change.status_code, 200)
        logout = client.post("/api/auth/logout")
        self.assertEqual(logout.status_code, 200)
        old = client.post("/api/auth/login", json={"username": "tester", "password": "secret123"})
        self.assertEqual(old.status_code, 401)
        new = client.post("/api/auth/login", json={"username": "tester", "password": "newsecret123"})
        self.assertEqual(new.status_code, 200)

    def test_07_outline_local_counts(self):
        for count in (6, 10, 18, 30, 50):
            outline = local_outline(f"测试主题 {count}", target_slides=count)
            self.assertEqual(len(outline["slides"]), count)
            self.assertEqual(outline["slides"][-1]["type"], "end")
            if count >= 30:
                self.assertGreaterEqual(sum(1 for slide in outline["slides"] if slide["type"] == "section"), 4)
            if count >= 40:
                agenda = next(slide for slide in outline["slides"] if slide["type"] == "agenda")
                self.assertEqual(len(agenda["bullets"]), 6)

    def test_08_admin_password_file_removed(self):
        from backend.config import ADMIN_CREDENTIALS_FILE

        client = self.app.test_client()
        self.assertEqual(
            client.post("/api/auth/login", json={"username": "admin", "password": "Admin@2016"}).status_code,
            200,
        )
        self.assertTrue(ADMIN_CREDENTIALS_FILE.exists())
        changed = client.post(
            "/api/admin/password",
            json={"current_password": "Admin@2016", "new_password": "AdminNew@2026"},
        )
        self.assertEqual(changed.status_code, 200, changed.get_json())
        self.assertFalse(ADMIN_CREDENTIALS_FILE.exists())

    def test_09_email_verification_code(self):
        import email
        from email import policy

        admin = self.app.test_client()
        self.assertEqual(
            admin.post("/api/auth/login", json={"username": "admin", "password": "AdminNew@2026"}).status_code,
            200,
        )
        configured = admin.put(
            "/api/admin/settings",
            json={
                "email_verification_enabled": True,
                "smtp": {
                    "host": "smtp.example.test",
                    "port": 465,
                    "security": "ssl",
                    "username": "",
                    "password": "",
                    "from_addr": "noreply@example.test",
                    "from_name": "启明 PPT 测试",
                },
            },
        )
        self.assertEqual(configured.status_code, 200, configured.get_json())

        client = self.app.test_client()
        sent = client.post(
            "/api/auth/send-code",
            json={"email": "mailuser@example.com", "invite_code": "songyonghao2016"},
        )
        self.assertEqual(sent.status_code, 200, sent.get_json())
        mail_dir = TEST_DATA / "mail"
        files = sorted(mail_dir.glob("*.eml"), key=lambda p: p.stat().st_mtime)
        self.assertTrue(files)
        message = email.message_from_bytes(files[-1].read_bytes(), policy=policy.default)
        body = message.get_body(preferencelist=("plain",)).get_content()
        import re

        match = re.search(r"\b(\d{6})\b", body)
        self.assertIsNotNone(match)
        code = match.group(1)

        wrong = client.post(
            "/api/auth/register",
            json={
                "username": "mailuser",
                "display_name": "邮箱用户",
                "email": "mailuser@example.com",
                "password": "mailpass123",
                "invite_code": "songyonghao2016",
                "email_code": "000000" if code != "000000" else "111111",
            },
        )
        self.assertEqual(wrong.status_code, 400)
        registered = client.post(
            "/api/auth/register",
            json={
                "username": "mailuser",
                "display_name": "邮箱用户",
                "email": "mailuser@example.com",
                "password": "mailpass123",
                "invite_code": "songyonghao2016",
                "email_code": code,
            },
        )
        self.assertEqual(registered.status_code, 200, registered.get_json())
        self.assertEqual(registered.get_json()["user"]["email"], "mailuser@example.com")

    def test_10_reference_uploads_and_outline(self):
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        from pptx.util import Inches
        from PIL import Image

        fixtures = TEST_DATA / "reference_fixtures"
        fixtures.mkdir(exist_ok=True)
        doc = Document()
        doc.add_heading("客户调研", 1)
        doc.add_paragraph("客户最关心交付速度，预算有限，需要先做小范围验证。")
        doc.save(fixtures / "brief.docx")
        wb = Workbook()
        ws = wb.active
        ws.title = "指标"
        ws.append(["指标", "数值"])
        ws.append(["活跃用户", 1200])
        ws.append(["转化率", "18%"])
        wb.save(fixtures / "metrics.xlsx")
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text_frame.text = "上一版结论：先聚焦一个高频场景。"
        prs.save(fixtures / "old_deck.pptx")
        (fixtures / "notes.txt").write_text("参考笔记：交付周期需要从 14 天缩短到 7 天。", encoding="utf-8")
        Image.new("RGB", (1200, 800), (30, 80, 150)).save(fixtures / "visual.png", "PNG")

        client = self.app.test_client()
        self.assertEqual(client.post("/api/auth/login", json={"username": "tester", "password": "newsecret123"}).status_code, 200)
        handles = {name: open(fixtures / name, "rb") for name in ("brief.docx", "metrics.xlsx", "old_deck.pptx", "notes.txt", "visual.png")}
        try:
            files = [(handle, name, "application/octet-stream") for name, handle in handles.items()]
            uploaded = client.post("/api/references/upload", data={"files": files}, content_type="multipart/form-data")
        finally:
            for handle in handles.values():
                handle.close()
        self.assertEqual(uploaded.status_code, 200, uploaded.get_json())
        payload = uploaded.get_json()
        self.assertEqual(len(payload["items"]), 5)
        self.assertFalse(payload["errors"])
        kinds = {item["filename"]: item["kind"] for item in payload["items"]}
        self.assertEqual(kinds["visual.png"], "image")
        self.assertEqual(kinds["brief.docx"], "document")
        ids = [item["id"] for item in payload["items"]]

        outline_response = client.post(
            "/api/ai/outline",
            json={"prompt": "结合参考资料做一份项目汇报", "slides": 8, "provider": "local", "reference_ids": ids},
        )
        self.assertEqual(outline_response.status_code, 200)
        outline = outline_response.get_json()["outline"]
        self.assertEqual(len(outline.get("reference_files") or []), 5)
        self.assertTrue(any("参考材料" in slide["title"] for slide in outline["slides"]))

        created = client.post(
            "/api/projects",
            json={
                "title": outline["title"],
                "prompt": "结合参考资料做一份项目汇报",
                "outline": outline,
                "template_id": "ink-paper",
                "options": {"animations": False, "audio": False, "video": False, "auto_advance": False, "reference_ids": ids, "use_reference_cover": True},
            },
        )
        self.assertEqual(created.status_code, 200, created.get_json())
        project = created.get_json()["project"]
        self.assertEqual(len(project.get("references") or []), 5)
        self.assertTrue(project["options"]["use_reference_cover"])

    def test_11_long_deck_generation(self):
        from backend.ai import local_outline
        from backend.pptx_builder import build_presentation, validate_presentation
        from backend.themes import get_theme

        outline = local_outline("AI 如何改变教育", target_slides=50, audience="中学老师")
        out_dir = TEST_DATA / "long_deck"
        out_dir.mkdir(exist_ok=True)
        output, meta = build_presentation(
            project_dir=out_dir,
            title=outline["title"],
            outline=outline,
            theme=get_theme("ink-paper"),
            options={"animations": True, "audio": False, "video": False, "auto_advance": False, "animation_level": "standard"},
        )
        self.assertTrue(output.exists())
        validation = validate_presentation(output)
        self.assertEqual(validation["slides"], 50)
        self.assertEqual(validation["animations"], 50)
        self.assertEqual(validation["transitions"], 50)
        self.assertEqual(meta["validation"]["slides"], 50)

    def test_12_media_autoplay_timing(self):
        from lxml import etree
        from backend.animations import build_animation_timing

        timing = build_animation_timing(
            targets=[(4, "fade", 500, 0)],
            media=[
                {"kind": "video", "shape_id": 2, "delay_ms": 0, "loop": True},
                {"kind": "audio", "shape_id": 12, "delay_ms": 650, "loop": False},
            ],
        )
        xml = etree.tostring(timing, encoding="unicode")
        self.assertIn('presetClass="mediacall"', xml)
        self.assertIn('cmd="playFrom(0.0)"', xml)
        self.assertIn("<p:audio", xml)
        self.assertIn("<p:video", xml)
        self.assertIn('delay="indefinite"', xml)


if __name__ == "__main__":
    unittest.main()
