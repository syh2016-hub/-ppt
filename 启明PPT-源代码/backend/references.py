"""Reference material ingestion: images, PDFs, Office documents and plain text."""
from __future__ import annotations

import csv
import io
import json
import mimetypes
import re
import shutil
import subprocess
import tempfile
import uuid
from pathlib import Path
from typing import Any

from PIL import Image, ImageOps

from .config import CACHE_DIR, DATA_DIR, REFERENCES_DIR
from .db import get_db, log_event
from .media import CREATE_NO_WINDOW, find_soffice
from .utils import clean_text, dominant_colors, human_size, now_iso, truncate

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_FILES = 8
MAX_TEXT_CHARS = 90_000
CONTEXT_CHARS = 18_000
TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".xml", ".html", ".htm",
    ".log", ".yaml", ".yml", ".ini", ".conf", ".py", ".js", ".ts", ".css", ".sql",
}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
PPTX_EXTS = {".pptx"}
XLSX_EXTS = {".xlsx", ".xlsm"}
LEGACY_EXTS = {
    ".doc", ".ppt", ".xls", ".rtf", ".odt", ".odp", ".ods", ".wps", ".wpd",
    ".pages", ".key", ".numbers", ".vsd", ".vsdx",
}
EXT_LABELS = {
    ".pdf": "PDF",
    ".docx": "Word",
    ".doc": "Word 97-2003",
    ".pptx": "PowerPoint",
    ".ppt": "PowerPoint 97-2003",
    ".xlsx": "Excel",
    ".xlsm": "Excel",
    ".xls": "Excel 97-2003",
    ".txt": "文本",
    ".md": "Markdown",
    ".csv": "CSV",
    ".tsv": "TSV",
    ".rtf": "RTF",
    ".odt": "ODT",
    ".odp": "ODP",
    ".ods": "ODS",
    ".html": "网页",
    ".htm": "网页",
}


def ext_label(ext: str) -> str:
    return EXT_LABELS.get(ext.lower(), (ext.lstrip(".") or "文件").upper())


def _safe_name(name: str) -> str:
    raw = str(name or "reference")
    base = re.split(r"[\\/]", raw)[-1]
    base = Path(base).name
    base = re.sub(r'[\x00-\x1f<>:"/\\|?*]+', "_", base).strip(" .")
    return (base or "reference")[:140]


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "gb18030", "big5", "latin1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _extract_pdf(path: Path) -> tuple[str, str]:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        chunks: list[str] = []
        for index, page in enumerate(reader.pages):
            if index >= 80:
                break
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            if text.strip():
                chunks.append(f"[第 {index + 1} 页]\n{text.strip()}")
        result = "\n\n".join(chunks)
        if not result.strip():
            return "", "PDF 中没有提取到文字，可能是扫描版；图片内容不会自动识别。"
        return result, ""
    except Exception as exc:
        return "", f"PDF 解析失败：{str(exc)[:120]}"


def _extract_docx(path: Path) -> tuple[str, str]:
    try:
        from docx import Document

        document = Document(str(path))
        chunks: list[str] = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables[:20]:
            for row in table.rows[:100]:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    chunks.append(" | ".join(cells))
        return "\n".join(chunks), ""
    except Exception as exc:
        return "", f"Word 文档解析失败：{str(exc)[:120]}"


def _extract_pptx(path: Path) -> tuple[str, str]:
    try:
        from pptx import Presentation

        presentation = Presentation(str(path))
        chunks: list[str] = []
        for index, slide in enumerate(presentation.slides):
            if index >= 60:
                break
            texts: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    value = shape.text_frame.text.strip()
                    if value:
                        texts.append(value)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    texts.append(f"备注：{notes}")
            if texts:
                chunks.append(f"[第 {index + 1} 页]\n" + "\n".join(texts))
        return "\n\n".join(chunks), ""
    except Exception as exc:
        return "", f"PPT 文档解析失败：{str(exc)[:120]}"


def _extract_xlsx(path: Path) -> tuple[str, str]:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(str(path), read_only=True, data_only=True)
        chunks: list[str] = []
        total_chars = 0
        for sheet in workbook.worksheets[:12]:
            chunks.append(f"[工作表：{sheet.title}]")
            for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_index >= 800 or total_chars > MAX_TEXT_CHARS:
                    break
                cells = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if cells:
                    line = " | ".join(cells)
                    chunks.append(line)
                    total_chars += len(line)
        workbook.close()
        return "\n".join(chunks), ""
    except Exception as exc:
        return "", f"Excel 表格解析失败：{str(exc)[:120]}"


def _extract_html(path: Path) -> tuple[str, str]:
    try:
        from bs4 import BeautifulSoup

        text = _decode_text(path.read_bytes())
        soup = BeautifulSoup(text, "lxml")
        for node in soup(["script", "style", "noscript"]):
            node.decompose()
        return soup.get_text("\n", strip=True), ""
    except Exception as exc:
        return "", f"网页解析失败：{str(exc)[:120]}"


def _extract_delimited(path: Path, delimiter: str) -> tuple[str, str]:
    try:
        text = _decode_text(path.read_bytes())
        rows = []
        for index, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter)):
            if index >= 1200:
                break
            cells = [str(cell).strip() for cell in row if str(cell).strip()]
            if cells:
                rows.append(" | ".join(cells))
        return "\n".join(rows), ""
    except Exception as exc:
        return "", f"表格解析失败：{str(exc)[:120]}"


def _extract_legacy(path: Path) -> tuple[str, str]:
    soffice = find_soffice()
    if not soffice:
        return "", "旧版/特殊格式需要本机安装 LibreOffice 才能提取正文。"
    output_dir = CACHE_DIR / "reference_convert" / uuid.uuid4().hex
    output_dir.mkdir(parents=True, exist_ok=True)
    profile = CACHE_DIR / "lo_reference_profile"
    profile.mkdir(parents=True, exist_ok=True)
    try:
        command = [
            str(soffice),
            "--headless", "--norestore", "--invisible",
            f"-env:UserInstallation=file:///{profile.as_posix()}",
            "--convert-to", "pdf", "--outdir", str(output_dir), str(path),
        ]
        subprocess.run(command, capture_output=True, timeout=90, creationflags=CREATE_NO_WINDOW)
        pdfs = list(output_dir.glob("*.pdf"))
        if not pdfs:
            return "", "LibreOffice 未能将这个文件转换为可读取的 PDF。"
        return _extract_pdf(pdfs[0])
    except Exception as exc:
        return "", f"旧版格式转换失败：{str(exc)[:120]}"
    finally:
        shutil.rmtree(output_dir, ignore_errors=True)


def extract_document_text(path: Path) -> tuple[str, str]:
    ext = path.suffix.lower()
    if ext in PDF_EXTS:
        text, warning = _extract_pdf(path)
    elif ext in DOCX_EXTS:
        text, warning = _extract_docx(path)
    elif ext in PPTX_EXTS:
        text, warning = _extract_pptx(path)
    elif ext in XLSX_EXTS:
        text, warning = _extract_xlsx(path)
    elif ext in {".html", ".htm"}:
        text, warning = _extract_html(path)
    elif ext == ".csv":
        text, warning = _extract_delimited(path, ",")
    elif ext == ".tsv":
        text, warning = _extract_delimited(path, "\t")
    elif ext in TEXT_EXTS:
        try:
            text, warning = _decode_text(path.read_bytes()), ""
        except Exception as exc:
            text, warning = "", f"文本读取失败：{str(exc)[:120]}"
    elif ext in LEGACY_EXTS:
        text, warning = _extract_legacy(path)
    else:
        raw = path.read_bytes()
        text = _decode_text(raw)
        printable = sum(1 for ch in text[:2000] if ch.isprintable() or ch in "\n\r\t")
        if raw and printable / max(1, len(text[:2000])) > 0.75:
            warning = "未识别的扩展名，已按纯文本读取。"
        else:
            return "", "暂不支持这种文件格式。请上传 PDF / Word / PPT / Excel / 图片 / 文本。"
    text = text.replace("\x00", "").strip()
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS] + "\n\n[内容过长，已截断]"
    return text, warning


def _image_palette(image: Image.Image) -> list[str]:
    try:
        return dominant_colors(image, 5)
    except Exception:
        return []


def _make_image(path: Path, directory: Path) -> tuple[Path, Path, int, int, list[str], str]:
    with Image.open(path) as source:
        image = ImageOps.exif_transpose(source)
        image.load()
        if getattr(image, "is_animated", False):
            image.seek(0)
        image = image.convert("RGB")
        width, height = image.size
        if max(width, height) > 2200:
            image.thumbnail((2200, 2200), Image.Resampling.LANCZOS)
        palette = _image_palette(image)
        normalized = directory / "image.jpg"
        image.save(normalized, "JPEG", quality=92, optimize=True)
        thumb = directory / "thumb.jpg"
        preview = image.copy()
        preview.thumbnail((720, 720), Image.Resampling.LANCZOS)
        preview.save(thumb, "JPEG", quality=84, optimize=True)
        warning = "" if width >= 500 and height >= 280 else "图片分辨率较低，用作封面时可能不够清晰。"
        return normalized, thumb, width, height, palette, warning


def _relative(path: Path) -> str:
    try:
        return path.resolve().relative_to(DATA_DIR.resolve()).as_posix()
    except ValueError:
        return path.as_posix()


def _absolute(relative: str) -> Path:
    value = (DATA_DIR / relative).resolve()
    if not str(value).startswith(str(DATA_DIR.resolve())):
        raise ValueError("参考文件路径异常")
    return value


def _insert_reference(
    user_id: int,
    *,
    filename: str,
    ext: str,
    kind: str,
    mime: str,
    path: Path,
    thumb_path: Path | None,
    text_path: Path | None,
    char_count: int,
    width: int,
    height: int,
    palette: list[str],
    size_bytes: int,
    warning: str,
) -> dict[str, Any]:
    db = get_db()
    cursor = db.execute(
        """INSERT INTO reference_files
           (user_id, filename, ext, kind, mime, path, thumb_path, text_path,
            char_count, width, height, palette_json, size_bytes, warning, created_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (
            user_id,
            filename,
            ext,
            kind,
            mime,
            _relative(path),
            _relative(thumb_path) if thumb_path else "",
            _relative(text_path) if text_path else "",
            char_count,
            width,
            height,
            json.dumps(palette, ensure_ascii=False),
            size_bytes,
            warning,
            now_iso(),
        ),
    )
    db.commit()
    row = db.execute("SELECT * FROM reference_files WHERE id = ?", (cursor.lastrowid,)).fetchone()
    return public_reference(dict(row))


def save_upload(user_id: int, upload) -> dict[str, Any]:
    filename = _safe_name(upload.filename or "reference")
    ext = Path(filename).suffix.lower()
    directory = REFERENCES_DIR / str(user_id) / uuid.uuid4().hex
    directory.mkdir(parents=True, exist_ok=True)
    original = directory / f"source{ext or '.bin'}"
    try:
        upload.save(original)
        size_bytes = original.stat().st_size
        if size_bytes <= 0:
            raise ValueError("文件为空")
        if size_bytes > MAX_FILE_BYTES:
            raise ValueError(f"单个文件不能超过 {human_size(MAX_FILE_BYTES)}")
        mime = upload.mimetype or mimetypes.guess_type(filename)[0] or "application/octet-stream"
        if ext in IMAGE_EXTS:
            try:
                normalized, thumb, width, height, palette, warning = _make_image(original, directory)
            except Exception as exc:
                raise ValueError(f"图片无法读取：{str(exc)[:100]}") from exc
            original.unlink(missing_ok=True)
            return _insert_reference(
                user_id,
                filename=filename,
                ext=ext,
                kind="image",
                mime="image/jpeg",
                path=normalized,
                thumb_path=thumb,
                text_path=None,
                char_count=0,
                width=width,
                height=height,
                palette=palette,
                size_bytes=size_bytes,
                warning=warning,
            )
        text, warning = extract_document_text(original)
        if not text.strip():
            raise ValueError(warning or "没有从文档中提取到可用文字")
        text_path = directory / "extracted.txt"
        text_path.write_text(text, encoding="utf-8")
        return _insert_reference(
            user_id,
            filename=filename,
            ext=ext,
            kind="document",
            mime=mime,
            path=original,
            thumb_path=None,
            text_path=text_path,
            char_count=len(text),
            width=0,
            height=0,
            palette=[],
            size_bytes=size_bytes,
            warning=warning,
        )
    except Exception:
        shutil.rmtree(directory, ignore_errors=True)
        raise


def public_reference(row: dict[str, Any]) -> dict[str, Any]:
    palette = []
    try:
        palette = json.loads(row.get("palette_json") or "[]")
    except ValueError:
        palette = []
    return {
        "id": int(row["id"]),
        "filename": row["filename"],
        "ext": row["ext"],
        "kind": row["kind"],
        "mime": row["mime"],
        "type_label": ext_label(row["ext"]),
        "char_count": int(row.get("char_count") or 0),
        "width": int(row.get("width") or 0),
        "height": int(row.get("height") or 0),
        "palette": palette,
        "size_bytes": int(row.get("size_bytes") or 0),
        "size_human": human_size(int(row.get("size_bytes") or 0)),
        "warning": row.get("warning") or "",
        "thumb_url": f"/api/references/{row['id']}/thumb" if row.get("thumb_path") else "",
        "created_at": row.get("created_at"),
    }


def get_reference(user_id: int, reference_id: int) -> dict[str, Any] | None:
    row = get_db().execute(
        "SELECT * FROM reference_files WHERE id = ? AND user_id = ?",
        (reference_id, user_id),
    ).fetchone()
    return dict(row) if row else None


def list_references(user_id: int, ids: list[int] | None = None) -> list[dict[str, Any]]:
    db = get_db()
    if ids:
        clean_ids = []
        for value in ids[:MAX_FILES]:
            try:
                clean_ids.append(int(value))
            except (TypeError, ValueError):
                continue
        if not clean_ids:
            return []
        placeholders = ",".join("?" for _ in clean_ids)
        rows = db.execute(
            f"SELECT * FROM reference_files WHERE user_id = ? AND id IN ({placeholders}) ORDER BY id DESC",
            (user_id, *clean_ids),
        ).fetchall()
    else:
        rows = db.execute(
            "SELECT * FROM reference_files WHERE user_id = ? ORDER BY id DESC LIMIT 60",
            (user_id,),
        ).fetchall()
    return [dict(row) for row in rows]


def delete_reference(user_id: int, reference_id: int) -> bool:
    row = get_reference(user_id, reference_id)
    if not row:
        return False
    db = get_db()
    db.execute("DELETE FROM reference_files WHERE id = ? AND user_id = ?", (reference_id, user_id))
    db.commit()
    source = _absolute(row["path"])
    directory = source.parent
    if str(directory.resolve()).startswith(str((REFERENCES_DIR / str(user_id)).resolve())):
        shutil.rmtree(directory, ignore_errors=True)
    return True


def thumb_path(user_id: int, reference_id: int) -> Path | None:
    row = get_reference(user_id, reference_id)
    if not row or not row.get("thumb_path"):
        return None
    target = _absolute(row["thumb_path"])
    return target if target.exists() else None


def reference_material(user_id: int, ids: list[int] | None, *, max_chars: int = CONTEXT_CHARS) -> dict[str, Any]:
    rows = list_references(user_id, ids or [])
    rows.reverse()
    chunks: list[str] = []
    files: list[dict[str, Any]] = []
    images: list[dict[str, Any]] = []
    warnings: list[str] = []
    used = 0
    for row in rows:
        public = public_reference(row)
        files.append(public)
        if row["kind"] == "image":
            images.append(
                {
                    "id": row["id"],
                    "filename": row["filename"],
                    "path": str(_absolute(row["path"])),
                    "palette": public["palette"],
                }
            )
            try:
                image_text = _absolute(row["text_path"]).read_text(encoding="utf-8") if row.get("text_path") else ""
            except OSError:
                image_text = ""
            if not image_text:
                image_text = _image_reading_text(row)
            if image_text.strip():
                remaining = max(0, max_chars - used)
                if remaining <= 0:
                    break
                excerpt = image_text.strip()[: min(1200, remaining)]
                used += len(excerpt)
                chunks.append(f"[图片《{row['filename']}》的 OCR 文字]\n{excerpt}")
            else:
                warnings.append(f"图片《{row['filename']}》没有识别到可读文字，仅作为视觉参考。")
            continue
        if not row.get("text_path"):
            continue
        try:
            text = _absolute(row["text_path"]).read_text(encoding="utf-8")
        except OSError:
            warnings.append(f"参考资料《{row['filename']}》读取失败")
            continue
        remaining = max(0, max_chars - used)
        if remaining <= 0:
            warnings.append("参考资料总长度超过上限，后续内容已截断。")
            break
        excerpt = text[:remaining]
        used += len(excerpt)
        chunks.append(f"《{row['filename']}》\n{excerpt}")
        if public.get("warning"):
            warnings.append(f"{row['filename']}：{public['warning']}")
    return {
        "text": "\n\n".join(chunks),
        "files": files,
        "images": images,
        "warnings": list(dict.fromkeys(warnings)),
    }


def reference_snippets(user_id: int, ids: list[int] | None, limit: int = 4) -> list[str]:
    material = reference_material(user_id, ids, max_chars=6000)
    snippets: list[str] = []
    for row in material["files"]:
        text = ""
        record = get_reference(user_id, int(row["id"]))
        if record and record.get("text_path"):
            try:
                text = _absolute(record["text_path"]).read_text(encoding="utf-8")
            except OSError:
                text = ""
        if not text:
            continue
        sentences = re.split(r"(?<=[。！？.!?])\s+|\n+", text)
        for sentence in sentences:
            sentence = re.sub(r"^\[[^\]]+\]\s*", "", sentence).strip()
            if 18 <= len(sentence) <= 160:
                snippets.append(sentence)
                break
        if len(snippets) >= limit:
            break
    return snippets


def first_image(user_id: int, ids: list[int] | None) -> Path | None:
    for row in list_references(user_id, ids or []):
        if row["kind"] == "image":
            target = _absolute(row["path"])
            if target.exists():
                return target
    return None


def validate_reference_ids(user_id: int, raw: Any) -> list[int]:
    if not isinstance(raw, list):
        return []
    result: list[int] = []
    for value in raw[:MAX_FILES]:
        try:
            reference_id = int(value)
        except (TypeError, ValueError):
            continue
        if get_reference(user_id, reference_id):
            result.append(reference_id)
    return list(dict.fromkeys(result))


_OCR_SCRIPT = r"""param([string]$Path)
$ErrorActionPreference = 'Stop'
[Console]::OutputEncoding = [System.Text.Encoding]::UTF8
Add-Type -AssemblyName System.Runtime.WindowsRuntime
$null = [Windows.Media.Ocr.OcrEngine, Windows.Foundation, ContentType=WindowsRuntime]
$null = [Windows.Storage.StorageFile, Windows.Storage, ContentType=WindowsRuntime]
$null = [Windows.Graphics.Imaging.BitmapDecoder, Windows.Graphics.Imaging, ContentType=WindowsRuntime]
$asTaskGeneric = ([System.WindowsRuntimeSystemExtensions].GetMethods() | Where-Object { $_.Name -eq 'AsTask' -and $_.GetParameters().Count -eq 1 -and $_.GetParameters()[0].ParameterType.Name -eq 'IAsyncOperation`1' })[0]
function Await($WinRtTask, $ResultType) {
  $asTask = $asTaskGeneric.MakeGenericMethod($ResultType)
  $netTask = $asTask.Invoke($null, @($WinRtTask))
  $netTask.Wait(-1) | Out-Null
  $netTask.Result
}
$file = Await ([Windows.Storage.StorageFile]::GetFileFromPathAsync($Path)) ([Windows.Storage.StorageFile])
$stream = Await ($file.OpenAsync([Windows.Storage.FileAccessMode]::Read)) ([Windows.Storage.Streams.IRandomAccessStream])
$decoder = Await ([Windows.Graphics.Imaging.BitmapDecoder]::CreateAsync($stream)) ([Windows.Graphics.Imaging.BitmapDecoder])
$bitmap = Await ($decoder.GetSoftwareBitmapAsync()) ([Windows.Graphics.Imaging.SoftwareBitmap])
$engine = [Windows.Media.Ocr.OcrEngine]::TryCreateFromUserProfileLanguages()
if (-not $engine) { Write-Output ''; exit 0 }
$result = Await ($engine.RecognizeAsync($bitmap)) ([Windows.Media.Ocr.OcrResult])
Write-Output $result.Text
"""


def _ocr_image(path: Path) -> str:
    """Read text from a textbook photo/page using Windows built-in OCR."""
    script = CACHE_DIR / "windows_ocr.ps1"
    if not script.exists():
        script.write_text(_OCR_SCRIPT, encoding="utf-8-sig")
    powershell = shutil.which("powershell") or shutil.which("pwsh")
    if not powershell:
        return ""
    try:
        result = subprocess.run(
            [powershell, "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script), "-Path", str(path)],
            capture_output=True,
            timeout=45,
            creationflags=CREATE_NO_WINDOW,
        )
        text = result.stdout.decode("utf-8", errors="ignore").strip()
        # Windows OCR inserts spaces between Chinese characters.
        text = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", text)
        text = re.sub(r"[ \t]{2,}", " ", text)
        return text.strip()
    except Exception:
        return ""


def _image_reading_text(row: dict[str, Any]) -> str:
    if row.get("text_path"):
        try:
            cached = _absolute(row["text_path"]).read_text(encoding="utf-8").strip()
            if cached:
                return cached
        except OSError:
            pass
    source = _absolute(row["path"])
    text = _ocr_image(source)
    if len(text) >= 20:
        try:
            text_path = source.parent / "ocr.txt"
            text_path.write_text(text, encoding="utf-8")
            db = get_db()
            db.execute(
                "UPDATE reference_files SET text_path = ?, char_count = ?, warning = ? WHERE id = ?",
                (_relative(text_path), len(text), "已通过 Windows OCR 识别图片文字。", row["id"]),
            )
            db.commit()
        except Exception:
            pass
    elif text:
        try:
            db = get_db()
            db.execute("UPDATE reference_files SET warning = ? WHERE id = ?", ("图片中文字较少或识别质量有限。", row["id"]))
            db.commit()
        except Exception:
            pass
    return text


def extract_reading_material(user_id: int, ids: list[int] | None) -> dict[str, Any]:
    """Find a textbook-like passage in documents or textbook images."""
    rows = list_references(user_id, ids or [])
    rows.reverse()
    image_count = sum(1 for row in rows if row.get("kind") == "image")
    ocr_count = 0
    for row in rows:
        if row.get("kind") == "document" and row.get("text_path"):
            try:
                raw = _absolute(row["text_path"]).read_text(encoding="utf-8")
            except OSError:
                raw = ""
            text = re.sub(r"\[第\s*\d+\s*页\]", "\n", raw)
            text = re.sub(r"\n{3,}", "\n\n", text).strip()
            paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
            label = re.compile(r"^(课文|朗读|阅读|课文朗读|Reading|Passage|Text)\s*[:：]?\s*", re.I)
            picked: list[str] = []
            for index, para in enumerate(paragraphs):
                if label.match(para) or re.search(r"(课文|朗读|Reading|Passage)", para[:40], re.I):
                    picked = [label.sub("", para)]
                    picked.extend(paragraphs[index + 1:index + 3])
                    break
            if not picked:
                candidates = [p for p in paragraphs if 60 <= len(p) <= 1800]
                if not candidates:
                    candidates = [p for p in paragraphs if len(p) >= 20]
                if candidates:
                    picked = [max(candidates[:8], key=len)]
            passage = "\n\n".join(picked).strip()
            passage = re.sub(r"^\s*(课文|朗读|阅读)\s*[:：]\s*", "", passage)
            if len(passage) >= 20:
                title = "课文朗读"
                first = paragraphs[0].splitlines()[0].strip() if paragraphs else ""
                if 2 <= len(first) <= 40 and not re.match(r"^\[?第\s*\d+", first):
                    title = first
                return {"title": title[:80], "text": passage[:3000], "source": row.get("filename", ""), "image_count": image_count, "ocr_count": ocr_count}
        elif row.get("kind") == "image":
            text = _image_reading_text(row)
            if len(text) >= 20:
                ocr_count += 1
                title_line = text.splitlines()[0].strip() if text.splitlines() else ""
                title = title_line[:40] if 2 <= len(title_line) <= 40 else "图片课文"
                return {"title": title, "text": text[:3000], "source": row.get("filename", ""), "image_count": image_count, "ocr_count": ocr_count}
    return {"title": "课文朗读", "text": "", "source": "", "image_count": image_count, "ocr_count": ocr_count}
